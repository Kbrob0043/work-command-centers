import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

OUTPUT_DIR = "/Users/knbed/.gemini/antigravity/scratch/work-command-center/downloads"
DOCS_DIR = "/Users/knbed/.gemini/antigravity/scratch/work-command-center/WORK COMMAND CENTER/Word Documents"
USER_DOWNLOADS = "/Users/knbed/Downloads"

os.makedirs(OUTPUT_DIR, exist_ok=True)
os.makedirs(DOCS_DIR, exist_ok=True)

PURPLE = RGBColor(92, 45, 145)      # #5c2d91
DARK_SLATE = RGBColor(15, 23, 42)    # #0f172a
GRAY_TEXT = RGBColor(100, 116, 139)  # #64748b
BLUE_LINK = RGBColor(37, 99, 235)    # #2563eb

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''<w:tcMar {nsdecls("w")}>
        <w:top w:w="{top}" w:type="dxa"/>
        <w:bottom w:w="{bottom}" w:type="dxa"/>
        <w:left w:w="{left}" w:type="dxa"/>
        <w:right w:w="{right}" w:type="dxa"/>
    </w:tcMar>''')
    tcPr.append(tcMar)

def add_portfolio_overview(doc):
    title = doc.add_paragraph()
    r = title.add_run("Portfolio Overview")
    r.font.name = "Segoe UI"
    r.font.size = Pt(24)
    r.font.bold = True
    r.font.color.rgb = DARK_SLATE
    title.paragraph_format.space_after = Pt(2)

    sub = doc.add_paragraph()
    r_sub = sub.add_run("K.O.  |  Friday, October 23, 2026")
    r_sub.font.name = "Segoe UI"
    r_sub.font.size = Pt(10)
    r_sub.font.color.rgb = GRAY_TEXT
    sub.paragraph_format.space_after = Pt(18)

    # Priorities
    h1 = doc.add_paragraph()
    r_h1 = h1.add_run("This Week’s Priorities")
    r_h1.font.name = "Segoe UI"
    r_h1.font.size = Pt(14)
    r_h1.font.bold = True
    r_h1.font.color.rgb = PURPLE
    h1.paragraph_format.space_after = Pt(6)

    for p in ["☐  Confirm ATMS+ UAT readiness", "☐  Resolve FOLAA+ scope questions", "☐  Capture GAL approval decisions"]:
        p_para = doc.add_paragraph()
        r_p = p_para.add_run(p)
        r_p.font.name = "Segoe UI"
        r_p.font.size = Pt(10.5)
        p_para.paragraph_format.space_after = Pt(3)

    # Health
    h2 = doc.add_paragraph()
    h2.paragraph_format.space_before = Pt(14)
    r_h2 = h2.add_run("Project Health")
    r_h2.font.name = "Segoe UI"
    r_h2.font.size = Pt(14)
    r_h2.font.bold = True
    r_h2.font.color.rgb = PURPLE
    h2.paragraph_format.space_after = Pt(6)

    table = doc.add_table(rows=4, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    col_widths = [Inches(1.8), Inches(1.5), Inches(2.2), Inches(1.2)]
    headers = ["Project", "Status", "Next milestone", "Due"]
    
    for i, h_text in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h_text
        set_cell_background(cell, "EDE7F6")
        set_cell_margins(cell)
        r = cell.paragraphs[0].runs[0]
        r.font.name = "Segoe UI"
        r.font.size = Pt(10)
        r.font.bold = True
        r.font.color.rgb = RGBColor(59, 7, 100)

    health_data = [
        ("ATMS+", "🔴 At risk", "UAT readiness", "Oct 28"),
        ("FOLAA+", "🟢 On track", "Scope review", "Oct 29"),
        ("GAL", "🟡 Waiting", "Sponsor approval", "Oct 30")
    ]
    for row_idx, data in enumerate(health_data, start=1):
        for col_idx, val in enumerate(data):
            cell = table.rows[row_idx].cells[col_idx]
            cell.text = val
            set_cell_margins(cell)
            r = cell.paragraphs[0].runs[0]
            r.font.name = "Segoe UI"
            r.font.size = Pt(10)
            if col_idx == 0:
                r.font.color.rgb = BLUE_LINK
                r.font.bold = True

    for row in table.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = width

    # Actions
    h3 = doc.add_paragraph()
    h3.paragraph_format.space_before = Pt(16)
    r_h3 = h3.add_run("Open Actions")
    r_h3.font.name = "Segoe UI"
    r_h3.font.size = Pt(14)
    r_h3.font.bold = True
    r_h3.font.color.rgb = PURPLE
    h3.paragraph_format.space_after = Pt(6)

    table2 = doc.add_table(rows=4, cols=4)
    table2.alignment = WD_TABLE_ALIGNMENT.CENTER
    table2.autofit = False
    col_widths2 = [Inches(2.8), Inches(1.4), Inches(1.1), Inches(1.4)]
    headers2 = ["Action", "Owner", "Due", "Status"]

    for i, h_text in enumerate(headers2):
        cell = table2.rows[0].cells[i]
        cell.text = h_text
        set_cell_background(cell, "EDE7F6")
        set_cell_margins(cell)
        r = cell.paragraphs[0].runs[0]
        r.font.name = "Segoe UI"
        r.font.size = Pt(10)
        r.font.bold = True
        r.font.color.rgb = RGBColor(59, 7, 100)

    actions_data = [
        ("Confirm test coverage", "K.O.", "Oct 27", "🟡 Open"),
        ("Review scope changes", "Project lead", "Oct 28", "⚫ Waiting"),
        ("Publish meeting decisions", "K.O.", "Oct 26", "🟡 Open")
    ]
    for row_idx, data in enumerate(actions_data, start=1):
        for col_idx, val in enumerate(data):
            cell = table2.rows[row_idx].cells[col_idx]
            cell.text = val
            set_cell_margins(cell)
            r = cell.paragraphs[0].runs[0]
            r.font.name = "Segoe UI"
            r.font.size = Pt(10)

    for row in table2.rows:
        for idx, width in enumerate(col_widths2):
            row.cells[idx].width = width

    # Decisions
    h4 = doc.add_paragraph()
    h4.paragraph_format.space_before = Pt(16)
    r_h4 = h4.add_run("Decisions & Escalations")
    r_h4.font.name = "Segoe UI"
    r_h4.font.size = Pt(14)
    r_h4.font.bold = True
    r_h4.font.color.rgb = PURPLE
    h4.paragraph_format.space_after = Pt(4)

    for d in ["• ATMS+: confirm release acceptance criteria", "• GAL: follow up on sponsor approval"]:
        p_dec = doc.add_paragraph()
        r_d = p_dec.add_run(d)
        r_d.font.name = "Segoe UI"
        r_d.font.size = Pt(10.5)
        p_dec.paragraph_format.space_after = Pt(3)

    # Quick Links
    h5 = doc.add_paragraph()
    h5.paragraph_format.space_before = Pt(16)
    r_h5 = h5.add_run("Quick Links")
    r_h5.font.name = "Segoe UI"
    r_h5.font.size = Pt(14)
    r_h5.font.bold = True
    r_h5.font.color.rgb = PURPLE
    h5.paragraph_format.space_after = Pt(4)

    links_p = doc.add_paragraph()
    r_links = links_p.add_run("Meeting template  •  Decision log  •  Risk register  •  Weekly review")
    r_links.font.name = "Segoe UI"
    r_links.font.size = Pt(10.5)
    r_links.font.color.rgb = BLUE_LINK

def add_this_week(doc):
    title = doc.add_paragraph()
    r = title.add_run("This Week")
    r.font.name = "Segoe UI"
    r.font.size = Pt(24)
    r.font.bold = True
    r.font.color.rgb = DARK_SLATE

    sub = doc.add_paragraph()
    r_sub = sub.add_run("Sprint & Daily Focus  |  Week of October 23, 2026")
    r_sub.font.name = "Segoe UI"
    r_sub.font.size = Pt(10)
    r_sub.font.color.rgb = GRAY_TEXT
    sub.paragraph_format.space_after = Pt(16)

    h1 = doc.add_paragraph()
    r_h1 = h1.add_run("Daily Schedule & Focus")
    r_h1.font.name = "Segoe UI"
    r_h1.font.size = Pt(14)
    r_h1.font.bold = True
    r_h1.font.color.rgb = PURPLE

    tbl = doc.add_table(rows=6, cols=3)
    tbl.autofit = False
    headers = ["Day", "Focus Area", "Key Meetings & Deliverables"]
    for i, h in enumerate(headers):
        cell = tbl.rows[0].cells[i]
        cell.text = h
        set_cell_background(cell, "EDE7F6")
        set_cell_margins(cell)
        r = cell.paragraphs[0].runs[0]
        r.font.name = "Segoe UI"
        r.font.size = Pt(10)
        r.font.bold = True
        r.font.color.rgb = RGBColor(59, 7, 100)

    schedule = [
        ("Monday", "Portfolio sync & ATMS+ kickoff", "9:00 AM Standup • 1:00 PM Core Team"),
        ("Tuesday", "FOLAA+ scope & Vendor review", "11:00 AM Vendor Review • 3:00 PM Workshop"),
        ("Wednesday", "GAL Sponsor approvals", "10:00 AM Sponsor Check-in • 2:00 PM CAB Gate"),
        ("Thursday", "Deep work & UAT prep", "UAT Test Pack Review • Deep work block"),
        ("Friday", "Weekly wrap-up & broadcast", "3:00 PM Retro • Publish executive summary")
    ]
    for row_idx, data in enumerate(schedule, start=1):
        for col_idx, val in enumerate(data):
            cell = tbl.rows[row_idx].cells[col_idx]
            cell.text = val
            set_cell_margins(cell)
            r = cell.paragraphs[0].runs[0]
            r.font.name = "Segoe UI"
            r.font.size = Pt(10)
            if col_idx == 0:
                r.font.bold = True
                r.font.color.rgb = PURPLE

    widths = [Inches(1.4), Inches(2.6), Inches(2.7)]
    for row in tbl.rows:
        for idx, w in enumerate(widths):
            row.cells[idx].width = w

# Create individual documents
doc_po = Document()
add_portfolio_overview(doc_po)
doc_po.save(os.path.join(DOCS_DIR, "Portfolio Overview.docx"))
doc_po.save(os.path.join(OUTPUT_DIR, "Portfolio Overview.docx"))
doc_po.save(os.path.join(USER_DOWNLOADS, "Portfolio Overview.docx"))

doc_tw = Document()
add_this_week(doc_tw)
doc_tw.save(os.path.join(DOCS_DIR, "This Week.docx"))
doc_tw.save(os.path.join(OUTPUT_DIR, "This Week.docx"))
doc_tw.save(os.path.join(USER_DOWNLOADS, "This Week.docx"))

# Create Master Notebook document containing all pages
doc_master = Document()
add_portfolio_overview(doc_master)
doc_master.add_page_break()
add_this_week(doc_master)
doc_master.save(os.path.join(DOCS_DIR, "WORK COMMAND CENTER (Master Notebook).docx"))
doc_master.save(os.path.join(OUTPUT_DIR, "WORK COMMAND CENTER (Master Notebook).docx"))
doc_master.save(os.path.join(USER_DOWNLOADS, "WORK COMMAND CENTER (Master Notebook).docx"))

print("All DOCX documents generated successfully!")
